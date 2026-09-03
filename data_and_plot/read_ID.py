from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# =============================================================================
# 参数设置
# =============================================================================

# 存放ID文件的子文件夹
INPUT_DIR = Path(r"./cosmos_1408/tle_history_output")

# 输出Excel文件
OUTPUT_XLSX = Path(r"./cosmos_1408_norad_ids.xlsx")

# 每行排列的列数
N_COLUMNS = 10

# 是否递归读取更深层的子文件夹
RECURSIVE = False

# 是否去掉文件扩展名
# True：49516.txt -> 49516
# False：49516.txt -> 49516.txt
REMOVE_EXTENSION = True

# 是否添加表头
ADD_HEADER = True

# =============================================================================
# 功能函数
# =============================================================================


def natural_sort_key(text: str):
    """按自然顺序排序，数字部分按照数值大小排序。"""
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", text)
    ]


def read_ids(folder: Path) -> list[str]:
    """读取文件夹中的文件名，并将文件名作为ID。"""
    if not folder.exists():
        raise FileNotFoundError(f"找不到输入文件夹：{folder.resolve()}")

    if not folder.is_dir():
        raise NotADirectoryError(f"输入路径不是文件夹：{folder.resolve()}")

    paths = folder.rglob("*") if RECURSIVE else folder.iterdir()

    ids = []
    for path in paths:
        if not path.is_file():
            continue

        file_id = path.stem if REMOVE_EXTENSION else path.name
        file_id = file_id.strip()

        if file_id:
            ids.append(file_id)

    # 去重后排序
    return sorted(set(ids), key=natural_sort_key)


def save_ids_to_excel(
    ids: list[str],
    output_path: Path,
    n_columns: int,
) -> None:
    """将排序后的ID按指定列数写入Excel。"""
    if n_columns <= 0:
        raise ValueError("N_COLUMNS必须是大于0的整数")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "NORAD IDs"

    data_start_row = 1

    if ADD_HEADER:
        for column in range(1, n_columns + 1):
            cell = worksheet.cell(
                row=1,
                column=column,
                value=f"NORAD ID",
            )
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        data_start_row = 2

    for index, file_id in enumerate(ids):
        row = index // n_columns + data_start_row
        column = index % n_columns + 1

        cell = worksheet.cell(
            row=row,
            column=column,
            value=file_id,
        )

        # 按文本储存，避免Excel自动改变ID
        cell.number_format = "@"
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for column in range(1, n_columns + 1):
        column_letter = get_column_letter(column)
        worksheet.column_dimensions[column_letter].width = 14

    worksheet.freeze_panes = ("A2" if ADD_HEADER else "A1")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    ids = read_ids(INPUT_DIR)

    if not ids:
        print(f"警告：文件夹中没有读取到文件：{INPUT_DIR.resolve()}")

    save_ids_to_excel(
        ids=ids,
        output_path=OUTPUT_XLSX,
        n_columns=N_COLUMNS,
    )

    data_rows = ((len(ids) + N_COLUMNS - 1) // N_COLUMNS if ids else 0)

    print(f"读取到的ID数量：{len(ids)}")
    print(f"每行列数：{N_COLUMNS}")
    print(f"数据行数：{data_rows}")
    print(f"输出文件：{OUTPUT_XLSX.resolve()}")


if __name__ == "__main__":
    main()
